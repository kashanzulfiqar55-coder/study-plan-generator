from __future__ import annotations

import io
import json
import os
import re
from pathlib import Path

import streamlit as st
from docx import Document
from pypdf import PdfReader

from workflow import StudyPackWorkflow, WorkflowConfig, WorkflowError

st.set_page_config(
    page_title="AI Study Pack Studio",
    page_icon="📚",
    layout="wide",
)

st.title("📚 AI Study Pack Studio")
st.write("A multi-stage AI workflow that plans, generates, assesses, reviews, and refines a personalized study pack.")
st.caption("Built with Python, Streamlit, and Groq. Verify important material against your syllabus.")

with st.sidebar:
    st.header("Workflow settings")
    secret_key = ""
    try:
        secret_key = st.secrets.get("GROQ_API_KEY", "")
    except Exception:
        pass
    api_key = st.text_input(
        "Groq API key",
        value=secret_key,
        type="password",
        help="For deployment, set GROQ_API_KEY in Streamlit Community Cloud → App settings → Secrets.",
    )
    model = st.text_input("Model", value="openai/gpt-oss-20b")
    retries = st.slider("Retries per stage", 0, 3, 2)
    show_stage_details = st.checkbox("Show workflow stage outputs", value=True)
    st.divider()
    st.markdown("**Workflow**")
    st.markdown("1. Plan\n2. Generate\n3. Assess\n4. Review\n5. Refine")

def extract_text(file) -> str:
    name = file.name.lower()
    data = file.getvalue()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="replace")
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(lines).strip()
    raise ValueError("Unsupported file type. Upload PDF, TXT, or DOCX.")

def safe_filename(value: str) -> str:
    result = re.sub(r"[^a-zA-Z0-9_-]+", "_", value.lower()).strip("_")
    return result or "study_pack"

def stage_label(stage: str) -> str:
    return {
        "Plan": "1/5 — Planning",
        "Generate": "2/5 — Content generation",
        "Assess": "3/5 — Quality assessment",
        "Review": "4/5 — Editorial review",
        "Refine": "5/5 — Refinement",
    }.get(stage, stage)

with st.form("study_pack_form"):
    left, right = st.columns(2)
    with left:
        subject = st.text_input("Subject", placeholder="e.g., Cyber Security")
        topic = st.text_input("Topic or chapter", placeholder="e.g., Network Security")
        learner_level = st.selectbox(
            "Learner level",
            [
                "Middle school",
                "Matric / O-Level",
                "Intermediate / A-Level",
                "University beginner",
                "University intermediate",
                "University advanced",
            ],
            index=3,
        )
        language = st.selectbox(
            "Output language",
            ["Simple English", "Urdu", "Roman Urdu", "English + Urdu"],
        )
    with right:
        difficulty = st.select_slider(
            "Quiz difficulty",
            options=["Easy", "Easy to medium", "Medium", "Medium to hard", "Hard"],
            value="Easy to medium",
        )
        question_count = st.slider("Quiz questions", 5, 30, 10, 5)
        requested_sections = st.multiselect(
            "Study pack sections",
            [
                "Chapter summary",
                "Key terms and definitions",
                "Step-by-step explanations",
                "Examples and analogies",
                "Flashcards",
                "Multiple-choice quiz",
                "Short-answer questions",
                "Answer key with explanations",
                "Common mistakes and exam tips",
                "Revision timetable",
            ],
            default=[
                "Chapter summary",
                "Key terms and definitions",
                "Step-by-step explanations",
                "Flashcards",
                "Multiple-choice quiz",
                "Answer key with explanations",
            ],
        )
    learning_goals = st.text_area(
        "Personal learning goals",
        placeholder="e.g., I am a beginner. Explain from basics and focus on concepts likely to appear in exams.",
        height=90,
    )
    uploaded_file = st.file_uploader("Optional source notes", type=["pdf", "txt", "docx"])
    pasted_notes = st.text_area(
        "Or paste notes here",
        placeholder="Paste class notes or the relevant chapter text...",
        height=150,
    )
    submitted = st.form_submit_button("Run 5-stage workflow", type="primary", use_container_width=True)

if submitted:
    if not api_key.strip():
        st.error("Enter a Groq API key in the sidebar, or configure GROQ_API_KEY in Streamlit secrets.")
        st.stop()
    if not requested_sections:
        st.error("Select at least one study pack section.")
        st.stop()
    if not subject.strip() and not topic.strip() and not pasted_notes.strip() and uploaded_file is None:
        st.error("Enter a subject/topic or provide source notes.")
        st.stop()

    source_text = pasted_notes.strip()
    if uploaded_file is not None:
        try:
            extracted = extract_text(uploaded_file)
            if extracted:
                source_text = (source_text + "\n\n" + extracted).strip()
            else:
                st.warning("No selectable text was extracted. Scanned PDFs need OCR.")
        except Exception as exc:
            st.error(f"Could not read the uploaded file: {exc}")
            st.stop()

    if len(source_text) > 45000:
        source_text = source_text[:45000]
        st.warning("Source notes exceeded 45,000 characters; the text was truncated for this run.")

    stage_status = {}
    status_area = st.status("Preparing workflow...", expanded=True)

    def on_stage(stage: str, status: str):
        stage_status[stage] = status
        status_area.write(f"**{stage_label(stage)}:** {status}")

    config = WorkflowConfig(
        model=model.strip(),
        max_retries=retries,
        max_tokens=7000,
        temperature=0.3,
    )

    try:
        status_area.update(label="Running multi-stage AI workflow...", state="running", expanded=True)
        workflow = StudyPackWorkflow(api_key, config, on_stage=on_stage)
        run = workflow.run(
            subject=subject,
            topic=topic,
            learner_level=learner_level,
            language=language,
            learning_goals=learning_goals,
            requested_sections=requested_sections,
            difficulty=difficulty,
            question_count=question_count,
            source_notes=source_text,
        )
        st.session_state["study_pack_run"] = {
            "final_pack": run.final_pack,
            "context": run.context,
            "stages": [
                {
                    "stage": stage.stage,
                    "content": stage.content,
                    "parsed_json": stage.parsed_json,
                    "elapsed_seconds": stage.elapsed_seconds,
                }
                for stage in run.stages
            ],
            "topic": topic or subject or "study_pack",
        }
        status_area.update(label="Workflow completed successfully", state="complete", expanded=False)
    except WorkflowError as exc:
        status_area.update(label="Workflow stopped", state="error", expanded=True)
        st.error(str(exc))
        st.info("Check your API key, selected model, account limits, and the stage error above. You can retry without changing your source notes.")
    except Exception as exc:
        status_area.update(label="Unexpected workflow error", state="error", expanded=True)
        st.error(f"Unexpected error: {exc}")
        st.info("Review the traceback in Streamlit logs. Do not expose API keys in logs or screenshots.")

if "study_pack_run" in st.session_state:
    result = st.session_state["study_pack_run"]
    st.divider()
    st.subheader("Final refined study pack")
    st.markdown(result["final_pack"])
    filename = safe_filename(result.get("topic", "study_pack"))
    st.download_button(
        "Download final pack (.md)",
        data=result["final_pack"].encode("utf-8"),
        file_name=f"{filename}_study_pack.md",
        mime="text/markdown",
        use_container_width=True,
    )
    st.download_button(
        "Download final pack (.txt)",
        data=result["final_pack"].encode("utf-8"),
        file_name=f"{filename}_study_pack.txt",
        mime="text/plain",
        use_container_width=True,
    )

    if show_stage_details:
        st.subheader("Workflow trace")
        st.caption("Each stage receives context from earlier stages. The plan, assessment, and review are passed to refinement.")
        for item in result["stages"]:
            title = stage_label(item["stage"])
            with st.expander(f"{title} · {item['elapsed_seconds']:.1f}s"):
                if item["parsed_json"] is not None:
                    st.json(item["parsed_json"])
                else:
                    st.markdown(item["content"])
        with st.expander("Context passed between stages"):
            safe_context = {
                key: value
                for key, value in result["context"].items()
                if key not in {"source_notes", "draft", "final_pack"}
            }
            st.json(safe_context)

st.divider()
st.caption("Privacy: uploaded notes are sent to Groq for generation when you run the workflow. Do not upload confidential or sensitive material.")
